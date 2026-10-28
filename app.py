import html
import logging
import time
from io import BytesIO
import streamlit.components.v1 as components
from pathlib import Path

import pandas as pd
import pyodbc as db
import streamlit as st
import plotly.graph_objects as go
from plotly.subplots import make_subplots

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)

# =========================================================
# 1. CONFIGURAÇÕES
# =========================================================

st.set_page_config(
    page_title="Ranking de Fornecedores | LCM",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="collapsed",
)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
)
logger = logging.getLogger("ranking_fornecedores")

COR_AZUL = "#17365D"
COR_VERDE = "#93C73E"
COR_VERDE_ESCURO = "#78AD25"
COR_VERDE_CLARO = "#F7FAF1"
COR_FUNDO = "#F3F5F8"
COR_TEXTO = "#494949"
COR_BORDA = "#D9E0E8"

MODELOS = {
    "empObr": {
        "nome": "Empresa e Obra",
        "arquivo": "ranking_empresa_obra",
    },
    "fornecedor": {
        "nome": "Por Fornecedor",
        "arquivo": "ranking_fornecedor",
    },
}

MESES_ANALISE = 12
CACHE_SEGUNDOS = 300
LIMITE_PREVIA = 5

# =========================================================
# 2. CONEXÃO COM O SQL SERVER
# =========================================================

DRIVER = "{ODBC Driver 17 for SQL Server}"


def conectar():
    """Abre a conexão usando as credenciais do secrets.toml."""
    return db.connect(
        f"DRIVER={DRIVER};"
        f"SERVER={st.secrets['uau']['server']};"
        f"DATABASE={st.secrets['uau']['database']};"
        f"UID={st.secrets['uau']['uid']};"
        f"PWD={st.secrets['uau']['pwd']};"
        "Connection Timeout=15;"
    )

# =========================================================
# OPÇÕES DOS FILTROS — CONSULTAS LEVES
# =========================================================

@st.cache_data(ttl=CACHE_SEGUNDOS, show_spinner=False)
def carregar_opcoes_empresas():
    sql = """
        SELECT DISTINCT
            E.Codigo_emp,
            CAST(E.Codigo_emp AS VARCHAR(20))
                + ' - ' + E.Desc_emp AS DescricaoEmpresa
        FROM ContasPagas CP
        INNER JOIN Empresas E
            ON CP.Empresa_Pag = E.Codigo_emp
        WHERE
            CP.DataEmissao_Pag >= DATEADD(MONTH, -12, GETDATE())
            AND CP.DataEmissao_Pag < DATEADD(
                DAY, 1, CAST(GETDATE() AS DATE)
            )
        ORDER BY DescricaoEmpresa
    """

    with conectar() as conexao:
        return pd.read_sql_query(sql, conexao)


@st.cache_data(ttl=CACHE_SEGUNDOS, show_spinner=False)
def carregar_opcoes_obras(empresas=()):
    sql = """
        SELECT DISTINCT
            O.cod_obr,
            O.cod_obr + ' - ' + O.descr_obr AS DescricaoObras
        FROM ContasPagas CP
        INNER JOIN Obras O
            ON CP.ObraProc_Pag = O.cod_obr
        WHERE
            CP.DataEmissao_Pag >= DATEADD(MONTH, -12, GETDATE())
            AND CP.DataEmissao_Pag < DATEADD(
                DAY, 1, CAST(GETDATE() AS DATE)
            )
    """

    parametros = []

    if empresas:
        marcadores = ",".join("?" for _ in empresas)
        sql += f" AND CP.Empresa_Pag IN ({marcadores})"
        parametros.extend(empresas)

    sql += " ORDER BY DescricaoObras"

    with conectar() as conexao:
        return pd.read_sql_query(
            sql,
            conexao,
            params=parametros,
        )


@st.cache_data(ttl=CACHE_SEGUNDOS, show_spinner=False)
def carregar_opcoes_fornecedores():
    sql = """
        SELECT DISTINCT
            CP.ChqNome_Pag
        FROM ContasPagas CP
        WHERE
            CP.DataEmissao_Pag >= DATEADD(MONTH, -12, GETDATE())
            AND CP.DataEmissao_Pag < DATEADD(
                DAY, 1, CAST(GETDATE() AS DATE)
            )
            AND CP.ChqNome_Pag IS NOT NULL
            AND LTRIM(RTRIM(CP.ChqNome_Pag)) <> ''
        ORDER BY CP.ChqNome_Pag
    """

    with conectar() as conexao:
        return pd.read_sql_query(sql, conexao)

# =========================================================
# 3. CONSULTA SQL E CARREGAMENTO
# =========================================================

SQL_BASE = """
    SELECT
        CP.Empresa_Pag,

        CAST(E.Codigo_emp AS VARCHAR(20))
            + ' - ' + E.Desc_emp AS DescricaoEmpresa,

        CP.ObraProc_Pag,

        O.cod_obr + ' - ' + O.descr_obr AS DescricaoObras,

        P.cod_pes,
        CP.ChqNome_Pag,

        CASE
            WHEN P.cpf_pes IS NULL THEN 'Sem Dados'

            WHEN LEN(P.cpf_pes) = 11 THEN
                STUFF(
                    STUFF(
                        STUFF(P.cpf_pes, 4, 0, '.'),
                        8, 0, '.'
                    ),
                    12, 0, '-'
                )

            WHEN LEN(P.cpf_pes) = 14 THEN
                STUFF(
                    STUFF(
                        STUFF(
                            STUFF(P.cpf_pes, 3, 0, '.'),
                            7, 0, '.'
                        ),
                        11, 0, '/'
                    ),
                    16, 0, '-'
                )

            ELSE P.cpf_pes
        END AS Identidade,

        SUM(CP.ValorProc_Pag) AS TotalRecebido,

        CAST(
            DATEADD(MONTH, -12, GETDATE()) AS DATE
        ) AS DataInicio,

        CAST(GETDATE() AS DATE) AS DataFim,

        SUM(SUM(CP.ValorProc_Pag)) OVER () AS TotalAno

    FROM ContasPagas CP

    LEFT JOIN Pessoas P
        ON CP.CodForn_Pag = P.cod_pes

    LEFT JOIN Obras O
        ON CP.ObraProc_Pag = O.cod_obr

    LEFT JOIN Empresas E
        ON CP.Empresa_Pag = E.Codigo_emp

    WHERE
        CP.DataEmissao_Pag >= DATEADD(
            MONTH, -12, GETDATE()
        )
        AND CP.DataEmissao_Pag < DATEADD(
            DAY, 1, CAST(GETDATE() AS DATE)
        )
"""

SQL_FINAL = """
    GROUP BY
        CP.Empresa_Pag,
        E.Codigo_emp,
        E.Desc_emp,
        CP.ObraProc_Pag,
        O.cod_obr,
        O.descr_obr,
        CP.ChqNome_Pag,
        P.cpf_pes,
        P.cod_pes

    ORDER BY TotalRecebido DESC
"""


@st.cache_data(ttl=CACHE_SEGUNDOS, show_spinner=False)
def carregar_total_geral():
    sql = """
        SELECT
            COALESCE(SUM(CP.ValorProc_Pag), 0) AS TotalGeral
        FROM ContasPagas CP
        WHERE
            CP.DataEmissao_Pag >= DATEADD(MONTH, -12, GETDATE())
            AND CP.DataEmissao_Pag < DATEADD(
                DAY, 1, CAST(GETDATE() AS DATE)
            )
    """

    conexao = conectar()

    try:
        cursor = conexao.cursor()
        try:
            cursor.execute(sql)
            resultado = cursor.fetchone()
            return float(resultado[0] or 0)
        finally:
            cursor.close()
    finally:
        conexao.close()

@st.cache_data(ttl=CACHE_SEGUNDOS, show_spinner=False)
def carregar_dados(empresas=(), obras=(), fornecedores=()):
    conexao = None
    cursor = None

    try:
        conexao = conectar()
        cursor = conexao.cursor()

        filtros_sql = []
        parametros = []

        if empresas:
            placeholders = ", ".join("?" for _ in empresas)
            filtros_sql.append(f"CP.Empresa_Pag IN ({placeholders})")
            parametros.extend(empresas)

        if obras:
            placeholders = ", ".join("?" for _ in obras)
            filtros_sql.append(f"CP.ObraProc_Pag IN ({placeholders})")
            parametros.extend(obras)

        if fornecedores:
            placeholders = ", ".join("?" for _ in fornecedores)
            filtros_sql.append(f"CP.ChqNome_Pag IN ({placeholders})")
            parametros.extend(fornecedores)

        # Só adiciona o AND quando existem filtros selecionados
        sql_filtros = ""
        if filtros_sql:
            sql_filtros = "\n AND " + "\n AND ".join(filtros_sql)

        consulta = SQL_BASE + sql_filtros + SQL_FINAL

        cursor.execute(consulta, parametros)
        colunas = [coluna[0] for coluna in cursor.description]
        registros = cursor.fetchall()

        return pd.DataFrame.from_records(
            registros,
            columns=colunas
        )

    except Exception:
        logging.exception("Erro ao consultar os pagamentos")
        raise

    finally:
        if cursor is not None:
            cursor.close()

        if conexao is not None:
            conexao.close()

# =========================================================
# 4. FORMATAÇÃO E PREPARAÇÃO DOS DADOS
# =========================================================

def assinatura_dataframe(df):
    hashes = pd.util.hash_pandas_object(
        df,
        index=True,
    )

    return (
        tuple(df.columns),
        int(hashes.sum()),
        len(df),
    )


def brl(valor):
    """Formata um valor monetário no padrão brasileiro."""
    if pd.isna(valor):
        valor = 0

    texto = f"{float(valor):,.2f}"
    return "R$ " + texto.replace(",", "X").replace(".", ",").replace("X", ".")


def texto_seguro(valor):
    if pd.isna(valor):
        return "Não informado"

    texto = str(valor).strip()
    return texto if texto else "Não informado"


def formatar_colunas_monetarias(df):
    """Cria uma cópia de apresentação sem modificar os números originais."""
    exibicao = df.copy()

    if "Total Pago" in exibicao.columns:
        exibicao["Total Pago"] = exibicao["Total Pago"].map(brl)

    return exibicao


def montar_modelo(df, modelo):
    """Agrupa os pagamentos conforme o modelo selecionado."""
    if modelo == "empObr":
        tabela = (
            df.groupby(
                [
                    "DescricaoEmpresa",
                    "DescricaoObras",
                    "ChqNome_Pag",
                ],
                as_index=False,
                dropna=False,
            )
            .agg(TotalPago=("TotalRecebido", "sum"))
            .rename(
                columns={
                    "DescricaoEmpresa": "Empresa",
                    "DescricaoObras": "Obra",
                    "ChqNome_Pag": "Fornecedor",
                    "TotalPago": "Total Pago",
                }
            )
        )

    
    elif modelo == "fornecedor":
        base = df[
            ["ChqNome_Pag", "DescricaoObras", "TotalRecebido"]
        ]

        # Calcula o total de cada fornecedor separadamente.
        totais = (
            base.groupby(
                "ChqNome_Pag",
                as_index=False,
                dropna=False,
                sort=False,
            )["TotalRecebido"]
            .sum()
            .rename(
                columns={
                    "ChqNome_Pag": "Fornecedor",
                    "TotalRecebido": "Total Pago",
                }
            )
        )

        # Remove obras repetidas antes de concatenar os nomes.
        obras_unicas = (
            base[["ChqNome_Pag", "DescricaoObras"]]
            .dropna(subset=["DescricaoObras"])
            .drop_duplicates()
            .sort_values(
                ["ChqNome_Pag", "DescricaoObras"],
                kind="stable",
            )
        )

        # Concatena as obras uma única vez por fornecedor.
        obras_agrupadas = (
            obras_unicas.groupby(
                "ChqNome_Pag",
                as_index=False,
                sort=False,
            )["DescricaoObras"]
            .agg(", ".join)
            .rename(
                columns={
                    "ChqNome_Pag": "Fornecedor",
                    "DescricaoObras": "Obras",
                }
            )
        )

        tabela = totais.merge(
            obras_agrupadas,
            on="Fornecedor",
            how="left",
            sort=False,
        )

        tabela["Obras"] = tabela["Obras"].fillna("")
        tabela = tabela[
            ["Fornecedor", "Obras", "Total Pago"]
        ]


    else:
        raise ValueError("O modelo selecionado é inválido.")

    tabela["Fornecedor"] = tabela["Fornecedor"].map(texto_seguro)

    return (
        tabela.sort_values(
            "Total Pago",
            ascending=False,
            kind="stable",
        )
        .reset_index(drop=True)
    )

# =========================================================
# 5. MÉTRICAS
# =========================================================

def calcular_metricas(df, df_filtrado):

    # Consulta o total geral sem aplicar filtros
    total_geral = carregar_total_geral()

    # Soma somente os dados filtrados
    total_filtrado = (
        float(df_filtrado["TotalRecebido"].sum())
        if not df_filtrado.empty
        else 0.0
    )

    qtd_empresas = df_filtrado["Empresa_Pag"].nunique()
    qtd_obras = df_filtrado["ObraProc_Pag"].nunique()
    qtd_fornecedores = df_filtrado["cod_pes"].nunique()

    return {
        "total_geral": total_geral,
        "total_filtrado": total_filtrado,
        "empresas": int(qtd_empresas),
        "obras": int(qtd_obras),
        "fornecedores": int(qtd_fornecedores),
    }

# =========================================================
# 6. GERAÇÃO DO EXCEL
# =========================================================

@st.cache_data(show_spinner=False)
def gerar_excel(df):
    inicio = time.perf_counter()
    arquivo = BytesIO()

    with pd.ExcelWriter(
        arquivo,
        engine="openpyxl",
    ) as writer:
        df.to_excel(
            writer,
            index=False,
            sheet_name="Fornecedores",
        )

        planilha = writer.sheets["Fornecedores"]
        planilha.freeze_panes = "A2"
        planilha.auto_filter.ref = planilha.dimensions

        from openpyxl.styles import (
            Alignment,
            Font,
            PatternFill,
        )
        from openpyxl.utils import get_column_letter

        for celula in planilha[1]:
            celula.font = Font(
                bold=True,
                color="FFFFFF",
            )
            celula.fill = PatternFill(
                fill_type="solid",
                fgColor="17365D",
            )
            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for coluna in planilha.columns:
            indice = coluna[0].column
            largura = max(
                len(str(celula.value or ""))
                for celula in coluna
            )
            planilha.column_dimensions[
                get_column_letter(indice)
            ].width = min(max(largura + 3, 12), 45)

        if "Total Pago" in df.columns:
            coluna_total = list(df.columns).index("Total Pago") + 1

            for linha in range(2, planilha.max_row + 1):
                planilha.cell(
                    row=linha,
                    column=coluna_total,
                ).number_format = '"R$" #,##0.00'

    logger.info(
        "Excel gerado em %.2f segundos.",
        time.perf_counter() - inicio,
    )
    return arquivo.getvalue()


# =========================================================
# 7. GERAÇÃO DO PDF
# =========================================================

@st.cache_data(show_spinner=False)
def gerar_pdf(df, nome_modelo, periodo):
    inicio = time.perf_counter()
    arquivo = BytesIO()

    documento = SimpleDocTemplate(
        arquivo,
        pagesize=landscape(A4),
        rightMargin=1.1 * cm,
        leftMargin=1.1 * cm,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
        title=f"Ranking de Fornecedores - {nome_modelo}",
    )

    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph(
            "Ranking de Fornecedores | LCM Construção",
            estilos["Title"],
        ),
        Paragraph(
            html.escape(nome_modelo),
            estilos["Heading2"],
        ),
        Paragraph(
            html.escape(periodo),
            estilos["Normal"],
        ),
        Spacer(1, 0.4 * cm),
    ]

    estilo_celula = ParagraphStyle(
        "CelulaRelatorio",
        parent=estilos["Normal"],
        fontName="Helvetica",
        fontSize=7,
        leading=9,
        wordWrap="CJK",
    )

    estilo_cabecalho = ParagraphStyle(
        "CabecalhoRelatorio",
        parent=estilo_celula,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )

    colunas = list(df.columns)
    dados = [[
        Paragraph(html.escape(str(coluna)), estilo_cabecalho)
        for coluna in colunas
    ]]

    for _, registro in df.iterrows():
        linha = []

        for coluna in colunas:
            valor = registro[coluna]

            if coluna == "Total Pago":
                conteudo = brl(valor)
            elif pd.isna(valor):
                conteudo = ""
            else:
                conteudo = str(valor)

            linha.append(
                Paragraph(
                    html.escape(conteudo),
                    estilo_celula,
                )
            )

        dados.append(linha)

    larguras = [None] * len(colunas)

    if len(colunas) == 4 and "Total Pago" in colunas:
        larguras = [
            4.5 * cm,
            6.0 * cm,
            10.0 * cm,
            3.0 * cm,
        ]
    elif len(colunas) == 3 and "Obras" in colunas:
        larguras = [
            5.0 * cm,
            15.5 * cm,
            3.0 * cm,
        ]

    tabela = Table(
        dados,
        colWidths=larguras,
        repeatRows=1,
        hAlign="LEFT",
    )

    tabela.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(COR_AZUL)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor(COR_BORDA)),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
                colors.white,
                colors.HexColor(COR_VERDE_CLARO),
            ]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ])
    )

    elementos.append(tabela)
    documento.build(elementos)

    logger.info(
        "PDF gerado em %.2f segundos.",
        time.perf_counter() - inicio,
    )
    return arquivo.getvalue()


# =========================================================
# 8. PRÉVIA HTML DO PDF
# =========================================================

def previa_pdf_html(df, modelo, periodo):
    nome_modelo = MODELOS[modelo]["nome"]
    previa = df.head(LIMITE_PREVIA)

    cabecalho = "".join(
        f"<th>{html.escape(str(coluna))}</th>"
        for coluna in previa.columns
    )

    linhas = []

    for _, registro in previa.iterrows():
        celulas = []

        for coluna in previa.columns:
            valor = registro[coluna]

            if pd.isna(valor):
                conteudo = ""

            elif coluna == "Total Pago":
                conteudo = html.escape(brl(valor))

            elif coluna == "Obras" and modelo == "fornecedor":
                obras = [
                    obra.strip()
                    for obra in str(valor).split(",")
                    if obra.strip()
                ]

                limite_obras = 3
                exibidas = obras[:limite_obras]

                conteudo = "<br>".join(
                    html.escape(obra)
                    for obra in exibidas
                )

                restantes = len(obras) - len(exibidas)

                if restantes > 0:
                    conteudo += (
                        '<div class="obras-restantes">'
                        f'+ {restantes} outras obras'
                        '</div>'
                    )

            else:
                conteudo = html.escape(str(valor))

            celulas.append(f"<td>{conteudo}</td>")

        linhas.append("<tr>" + "".join(celulas) + "</tr>")

    corpo = "".join(linhas)

    if previa.empty:
        corpo = (
            '<tr><td colspan="10">'
            'Nenhum registro disponível.'
            '</td></tr>'
        )

    return f"""
        <style>
            body {{
                margin: 0;
                font-family: Arial, sans-serif;
                color: {"#F1F5F9" if tema_escuro else "#494949"};
                background: {"#222830" if tema_escuro else "#FFFFFF"};
            }}

            .previa-pdf {{
                border: 1px solid {"#3B434D" if tema_escuro else "#D9E0E8"};
                border-radius: 10px;
                padding: 12px;
                box-sizing: border-box;
            }}

            .previa-pdf-titulo {{
                font-size: 13px;
                font-weight: 700;
                margin-bottom: 5px;
            }}

            .previa-pdf-sub {{
                font-size: 10px;
                color: {"#A7B0BC" if tema_escuro else "#64748B"};
                margin-bottom: 12px;
            }}

            .previa-pdf table {{
                width: 100%;
                table-layout: fixed;
                border-collapse: collapse;
                font-size: 10px;
            }}

            .previa-pdf th {{
                background: #17365D;
                color: white;
                text-align: left;
                padding: 7px;
            }}

            .previa-pdf td {{
                padding: 7px;
                border-bottom: 1px solid {"#3B434D" if tema_escuro else "#D9E0E8"};
                vertical-align: top;
                overflow-wrap: anywhere;
            }}

            .previa-pdf tr:nth-child(even) td {{
                background: {"#29351C" if tema_escuro else "#F7FAF1"};
            }}

            .obras-restantes {{
                margin-top: 5px;
                font-size: 9px;
                font-weight: 700;
                color: {"#A7B0BC" if tema_escuro else "#64748B"};
            }}
        </style>

        <div class="previa-pdf">
            <div class="previa-pdf-titulo">
                {html.escape(nome_modelo)}
            </div>

            <div class="previa-pdf-sub">
                {html.escape(periodo)}
            </div>

            <table>
                <thead>
                    <tr>{cabecalho}</tr>
                </thead>
                <tbody>{corpo}</tbody>
            </table>
        </div>
    """


# =========================================================
# 9. TEMAS E CSS
# =========================================================

if "modelo_selecionado" not in st.session_state:
    st.session_state.modelo_selecionado = "empObr"

if "tema_escuro" not in st.session_state:
    st.session_state.tema_escuro = False

tema_escuro = st.session_state.tema_escuro

if tema_escuro:
    variaveis = """
        --fundo: #171B20;
        --superficie: #222830;
        --superficie-secundaria: #292F38;
        --texto: #F1F5F9;
        --texto-secundario: #A7B0BC;
        --borda: #3B434D;
        --verde-claro: #29351C;
        --campo: #20262D;
    """
else:
    variaveis = """
        --fundo: #F3F5F8;
        --superficie: #FFFFFF;
        --superficie-secundaria: #F7FAF1;
        --texto: #494949;
        --texto-secundario: #64748B;
        --borda: #D9E0E8;
        --verde-claro: #F7FAF1;
        --campo: #FFFFFF;
    """

st.markdown(
    f"""
    <style>
    :root {{
        --azul: {COR_AZUL};
        --verde: {COR_VERDE};
        --verde-escuro: {COR_VERDE_ESCURO};
        {variaveis}
        --raio: 14px;
    }}

    .stApp {{
        background: var(--fundo);
        color: var(--texto);
    }}

    [data-testid="stHeader"] {{
        background: transparent;
    }}

    [data-testid="stMainBlockContainer"] {{
        padding-top: 1.4rem;
        padding-bottom: 3rem;
        max-width: 1500px;
    }}

    h1, h2, h3, p, label,
    [data-testid="stMarkdownContainer"] {{
        color: var(--texto);
    }}

    .titulo-pagina {{
        font-size: clamp(26px, 3vw, 34px);
        line-height: 1.2;
        font-weight: 800;
        letter-spacing: -0.7px;
        margin-bottom: 5px;
        color: var(--texto);
    }}

    .descricao-pagina {{
        font-size: 14px;
        color: var(--texto-secundario);
        margin-bottom: 20px;
    }}

    .subtitulo-pagina {{
        font-size: 18px;
        font-weight: 750;
        color: var(--texto);
        margin: 22px 0 10px 0;
    }}

    .card-metrica {{
        background: var(--superficie);
        border: 1px solid var(--borda);
        border-top: 3px solid var(--verde);
        border-radius: var(--raio);
        padding: 16px;
        min-height: 112px;
        margin-bottom: 8px;
        overflow-wrap: anywhere;
    }}

    .card-metrica.destaque {{
        background: var(--verde);
        border-color: var(--verde);
        border-top: 3px solid var(--verde-escuro);
    }}

    .card-metrica .rotulo {{
        color: var(--texto-secundario);
        font-size: 11px;
        font-weight: 750;
        letter-spacing: .04em;
        text-transform: uppercase;
    }}

    .card-metrica .valor {{
        color: var(--texto);
        font-size: clamp(20px, 2vw, 27px);
        font-weight: 800;
        margin-top: 9px;
        line-height: 1.25;
    }}

    .card-metrica.destaque .rotulo,
    .card-metrica.destaque .detalhe {{
        color: #E1E8F0;
    }}

    .card-metrica.destaque .valor {{
        color: #FFFFFF;
    }}

    .card-metrica .detalhe {{
        color: var(--texto-secundario);
        font-size: 11px;
        margin-top: 6px;
    }}

    .modelo-titulo {{
        color: var(--texto);
        font-size: 15px;
        font-weight: 750;
        margin-bottom: 8px;
    }}

    .previa-pdf {{
        background: var(--superficie);
        color: var(--texto);
        border: 1px solid var(--borda);
        border-radius: 10px;
        padding: 12px;
        min-height: 215px;
        overflow-x: auto;
    }}

    .previa-pdf-titulo {{
        color: var(--texto);
        font-size: 13px;
        font-weight: 750;
    }}

    .previa-pdf-sub {{
        color: var(--texto-secundario);
        font-size: 10px;
        margin: 5px 0 12px;
    }}

    .previa-pdf table {{
        width: 100%;
        border-collapse: collapse;
        font-size: 10px;
    }}

    .previa-pdf th {{
        background: var(--azul);
        color: #FFFFFF;
        text-align: left;
        padding: 7px;
    }}

    .previa-pdf td {{
        padding: 7px;
        border-bottom: 1px solid var(--borda);
        color: var(--texto);
        overflow-wrap: anywhere;
    }}

    .previa-pdf tr:nth-child(even) td {{
        background: var(--verde-claro);
    }}

    div[data-testid="stVerticalBlockBorderWrapper"] {{
        border-color: var(--borda);
        border-radius: var(--raio);
    }}

    div[data-testid="stDownloadButton"] button,
    div[data-testid="stButton"] button {{
        min-height: 44px;
        border-radius: 10px;
        font-weight: 700;
        transition: background .15s ease, border-color .15s ease;
    }}

    div[data-testid="stDownloadButton"] button {{
        background: var(--verde) !important;
        color: #17200B !important;
        border: 1px solid var(--verde) !important;
    }}

    div[data-testid="stDownloadButton"] button:hover {{
        background: var(--verde-escuro) !important;
        color: #FFFFFF !important;
        border-color: var(--verde-escuro) !important;
    }}

    div[data-testid="stDownloadButton"] button p {{
        color: inherit !important;
    }}

    div[data-testid="stMultiSelect"] > div > div,
    div[data-testid="stSelectbox"] > div > div {{
        background: var(--campo);
        border-color: var(--borda);
    }}

    div[data-baseweb="select"] span {{
        color: var(--texto);
    }}

    div[data-testid="stDataFrame"] {{
        border: 1px solid var(--borda);
        border-radius: 10px;
        overflow: hidden;
    }}

    button:focus-visible,
    input:focus-visible {{
        outline: 2px solid var(--verde) !important;
        outline-offset: 2px !important;
    }}

    @media (max-width: 700px) {{
        .card-metrica {{
            padding: 12px;
            min-height: 98px;
        }}
        .card-metrica .valor {{
            font-size: 20px;
        }}
    }}
    .previa-pdf table {{
        width: 100%;
        table-layout: fixed;
        border-collapse: collapse;
        font-size: 10px;
    }}

    .previa-pdf th,
    .previa-pdf td {{
        padding: 7px;
        vertical-align: top;
        overflow-wrap: anywhere;
        word-break: normal;
    }}

    .obras-restantes {{
        margin-top: 5px;
        color: var(--texto-secundario);
        font-weight: 700;
        font-size: 9px;
    }}
    /* Botão Buscar dados */
    div[data-testid="stButton"]:has(button[kind="primary"])
    button {{
        background-color: var(--verde) !important;
        color: #17200B !important;
        border: 1px solid var(--verde) !important;
        font-weight: 700 !important;
        border-radius: 10px !important;
        min-height: 44px !important;
    }}

    /* Efeito ao passar o mouse */
    div[data-testid="stButton"]:has(button[kind="primary"])
    button:hover {{
        background-color: var(--verde-escuro) !important;
        color: #FFFFFF !important;
        border-color: var(--verde-escuro) !important;
    }}
    </style>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# 16. CURVA ABC
# =========================================================

CORES_ABC = {
    "A": "#93C73E",
    "B": "#F97316",
    "C": "#64748B",
}


def calcular_curva_abc(df):
    base = df[["ChqNome_Pag", "TotalRecebido"]].copy()
    base = base.dropna(subset=["ChqNome_Pag", "TotalRecebido"])
    base["ChqNome_Pag"] = base["ChqNome_Pag"].astype(str).str.strip()
    base = base[base["ChqNome_Pag"] != ""]

    abc = (
        base.groupby("ChqNome_Pag", as_index=False)["TotalRecebido"]
        .sum()
        .rename(
            columns={
                "ChqNome_Pag": "Fornecedor",
                "TotalRecebido": "Total Pago",
            }
        )
        .sort_values("Total Pago", ascending=False)
        .reset_index(drop=True)
    )

    total = abc["Total Pago"].sum()

    if abc.empty or total <= 0:
        return pd.DataFrame()

    abc["Participação (%)"] = abc["Total Pago"] / total * 100
    abc["Acumulado (%)"] = abc["Participação (%)"].cumsum()

    acumulado_anterior = abc["Acumulado (%)"] - abc["Participação (%)"]

    abc["Classe"] = "C"
    abc.loc[acumulado_anterior < 95, "Classe"] = "B"
    abc.loc[acumulado_anterior < 80, "Classe"] = "A"

    return abc


def _resumo_por_classe(abc):
    total_valor = abc["Total Pago"].sum()
    total_forn = len(abc)
    linhas = []
    for classe in ["A", "B", "C"]:
        parte = abc[abc["Classe"] == classe]
        linhas.append(
            {
                "Classe": classe,
                "Fornecedores": len(parte),
                "% dos fornecedores": len(parte) / total_forn * 100,
                "Valor pago": parte["Total Pago"].sum(),
                "% do valor": parte["Total Pago"].sum() / total_valor * 100,
            }
        )
    return pd.DataFrame(linhas)


def _rotulo_curto(nome, tamanho=22):
    nome = str(nome)
    return nome if len(nome) <= tamanho else nome[: tamanho - 1] + "…"


def _titulo_secao(texto):
    st.markdown(
        f'<div class="subtitulo-pagina">{html.escape(texto)}</div>',
        unsafe_allow_html=True,
    )


def exibir_grafico_abc(df_filtrado):
    abc = calcular_curva_abc(df_filtrado)

    if abc.empty:
        st.info("Não há valores pagos suficientes para gerar a Curva ABC.")
        return

    total_geral = abc["Total Pago"].sum()
    resumo = _resumo_por_classe(abc)

    # ---------- Quantidade de fornecedores exibidos ----------
    if len(abc) > 10:
        limite = st.slider(
            "Fornecedores exibidos no gráfico",
            min_value=10,
            max_value=min(100, len(abc)),
            value=min(30, len(abc)),
            step=5 if len(abc) > 20 else 1,
            key="abc_limite",
        )
    else:
        limite = len(abc)

    grafico = abc.head(limite).copy()
    n = len(grafico)
    grafico["Posição"] = range(1, n + 1)
    grafico["Rótulo"] = grafico["Fornecedor"].map(_rotulo_curto)

    cor_linha = "#FFFFFF" if tema_escuro else "#17365D"
    cor_texto = "#E2E8F0" if tema_escuro else "#334155"
    cor_grade = "rgba(148,163,184,0.18)"

    figura = make_subplots(specs=[[{"secondary_y": True}]])

    # ---------- Faixas de fundo A / B / C ----------
    for y0, y1, classe in [(0, 80, "A"), (80, 95, "B"), (95, 105, "C")]:
        figura.add_hrect(
            y0=y0, y1=y1, fillcolor=CORES_ABC[classe], opacity=0.07,
            line_width=0, layer="below", secondary_y=True,
        )

    # ---------- Barras: uma trace por classe ----------
    for classe in ["A", "B", "C"]:
        parte = grafico[grafico["Classe"] == classe]
        if parte.empty:
            continue

        figura.add_trace(
            go.Bar(
                x=parte["Posição"],
                y=parte["Total Pago"],
                name=f"Classe {classe}",
                marker=dict(color=CORES_ABC[classe], line=dict(width=0)),
                opacity=0.92,
                customdata=parte[
                    ["Fornecedor", "Participação (%)", "Acumulado (%)"]
                ].to_numpy(),
                hovertemplate=(
                    "<b>%{customdata[0]}</b><br>"
                    f"Classe {classe} · posição %{{x}}<br>"
                    "Total pago: R$ %{y:,.2f}<br>"
                    "Participação: %{customdata[1]:.2f}%<br>"
                    "Acumulado: %{customdata[2]:.2f}%"
                    "<extra></extra>"
                ),
            ),
            secondary_y=False,
        )

    # ---------- Linha do acumulado ----------
    figura.add_trace(
        go.Scatter(
            x=[0] + grafico["Posição"].tolist(),
            y=[0] + grafico["Acumulado (%)"].tolist(),
            name="% acumulado",
            mode="lines",
            line=dict(color=cor_linha, width=3, shape="spline", smoothing=0.6),
            hovertemplate="Acumulado: %{y:.2f}%<extra></extra>",
        ),
        secondary_y=True,
    )

    # Marcadores apenas onde a curva cruza 80% e 95%
    for limite_classe, rotulo in [(80, "A"), (95, "B")]:
        passou = grafico[grafico["Acumulado (%)"] >= limite_classe]
        if passou.empty:
            continue
        ponto = passou.iloc[0]
        figura.add_trace(
            go.Scatter(
                x=[ponto["Posição"]],
                y=[ponto["Acumulado (%)"]],
                mode="markers+text",
                marker=dict(
                    size=11,
                    color=CORES_ABC[rotulo],
                    line=dict(color=cor_linha, width=2),
                ),
                text=[f"{int(ponto['Posição'])} fornec."],
                textposition="top left",
                textfont=dict(color=cor_texto, size=12),
                showlegend=False,
                hoverinfo="skip",
                cliponaxis=False,
            ),
            secondary_y=True,
        )

    # ---------- Linhas de referência ----------
    figura.add_hline(
        y=80, line_dash="dot", line_color=CORES_ABC["A"], line_width=1.5,
        annotation_text="80%", annotation_position="right",
        annotation_font_color=CORES_ABC["A"], secondary_y=True,
    )
    figura.add_hline(
        y=95, line_dash="dot", line_color=CORES_ABC["B"], line_width=1.5,
        annotation_text="95%", annotation_position="right",
        annotation_font_color=CORES_ABC["B"], secondary_y=True,
    )

    # ---------- Layout ----------
    figura.update_layout(
        title=dict(
            text=(
                "Curva ABC de Fornecedores"
                f"<br><sup>Top {n} de {len(abc)} fornecedores · "
                f"{grafico['Total Pago'].sum() / total_geral * 100:.1f}% "
                "do valor total pago</sup>"
            ),
            x=0,
            xanchor="left",
        ),
        template="plotly_dark" if tema_escuro else "plotly_white",
        height=600,
        bargap=0.25,
        hovermode="closest",
        font=dict(color=cor_texto),
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        margin=dict(l=20, r=50, t=100, b=30),
        legend=dict(
            orientation="h",
            yanchor="bottom",
            y=1.02,
            xanchor="right",
            x=1,
        ),
        hoverlabel=dict(font_size=13),
    )

    # Poucos fornecedores: nomes. Muitos: ranking (nome no hover).
    if n <= 20:
        figura.update_xaxes(
            tickmode="array",
            tickvals=grafico["Posição"].tolist(),
            ticktext=grafico["Rótulo"].tolist(),
            tickangle=-40,
            automargin=True,
            title_text=None,
        )
    else:
        figura.update_xaxes(
            tickmode="array",
            tickvals=[1] + list(range(5, n + 1, 5)),
            automargin=True,
            title_text="Ranking do fornecedor (passe o mouse para ver o nome)",
        )

    figura.update_xaxes(
        range=[0, n + 0.6],
        showgrid=False,
        zeroline=False,
    )

    figura.update_yaxes(
        title_text="Total pago",
        tickprefix="R$ ",
        tickformat="~s",
        rangemode="tozero",
        showgrid=True,
        gridcolor=cor_grade,
        zeroline=False,
        automargin=True,
        secondary_y=False,
    )

    figura.update_yaxes(
        title_text="Percentual acumulado",
        ticksuffix="%",
        range=[0, 105],
        tickvals=[0, 20, 40, 60, 80, 95, 100],
        showgrid=False,
        zeroline=False,
        automargin=True,
        secondary_y=True,
    )

    with st.container(border=True):
        st.plotly_chart(
            figura,
            use_container_width=True,
            config={"displaylogo": False},
        )

    # ---------- Resumo por classe ----------
    _titulo_secao("Resumo por classe")

    resumo_fmt = resumo.copy()
    resumo_fmt["% dos fornecedores"] = resumo_fmt["% dos fornecedores"].map(
        lambda v: f"{v:.1f}%"
    )
    resumo_fmt["% do valor"] = resumo_fmt["% do valor"].map(
        lambda v: f"{v:.1f}%"
    )
    resumo_fmt["Valor pago"] = resumo_fmt["Valor pago"].map(brl)

    st.dataframe(resumo_fmt, use_container_width=True, hide_index=True)

    # ---------- Tabela completa ----------
    _titulo_secao("Classificação dos fornecedores")

    tabela = abc.copy()
    tabela.insert(0, "#", range(1, len(tabela) + 1))

    st.dataframe(
        tabela.head(500),
        use_container_width=True,
        hide_index=True,
        height=400,
        column_config={
            "Total Pago": st.column_config.NumberColumn(
                "Total Pago", format="R$ %.2f"
            ),
            "Participação (%)": st.column_config.ProgressColumn(
                "Participação (%)",
                format="%.2f%%",
                min_value=0,
                max_value=float(tabela["Participação (%)"].max()),
            ),
            "Acumulado (%)": st.column_config.NumberColumn(
                "Acumulado (%)", format="%.2f%%"
            ),
        },
    )

    if len(tabela) > 500:
        st.caption(
            "A tabela exibe os primeiros 500 fornecedores. "
            "A classificação considera todos os fornecedores."
        )



aba_ranking, aba_relatorios, aba_abc = st.tabs([
    "📊 Ranking", "📁 Relatórios", "📈 Curva ABC"
])

with aba_ranking:
    # =========================================================
    # 10. CABEÇALHO E SELEÇÃO DO MODELO
    # =========================================================

    cabecalho_col, tema_col = st.columns([5, 1])

    with cabecalho_col:
        st.markdown(
            '<div class="titulo-pagina">Ranking de Fornecedores</div>'
            '<div class="descricao-pagina">'
            'Analise e consolide os pagamentos por empresa, obra e fornecedor.'
            '</div>',
            unsafe_allow_html=True,
        )

    with tema_col:
        novo_tema = st.toggle(
            "Tema escuro",
            value=st.session_state.tema_escuro,
            key="controle_tema",
        )

        if novo_tema != st.session_state.tema_escuro:
            st.session_state.tema_escuro = novo_tema
            st.rerun()

    from streamlit.components.v1 import declare_component

    CAMINHO_COMPONENTE = Path(__file__).resolve().parent / "selecionarCard"
    CAMINHO_HTML = CAMINHO_COMPONENTE / "index.html"

    selecionarCard = declare_component(
        "selecionarCard",
        path=str(CAMINHO_COMPONENTE),
    )

    st.markdown(
        '<div class="subtitulo-pagina">Escolha o modelo</div>',
        unsafe_allow_html=True,
    )

    valor_modelo = selecionarCard(
        valor_inicial=st.session_state.modelo_selecionado,
        tema_escuro=st.session_state.tema_escuro,
        key="seletor_modelos",
        default=st.session_state.modelo_selecionado,
    )

    if valor_modelo in MODELOS:
        if valor_modelo != st.session_state.modelo_selecionado:
            st.session_state.modelo_selecionado = valor_modelo
            for chave in ("busca_executada", "filtros_confirmados", "dados_ranking", "tabela_ranking", "periodo_ranking", "modelo_ranking", "arquivo_excel", "arquivo_pdf", "chave_excel", "chave_pdf"):
                st.session_state.pop(chave, None)
            st.session_state.pop("filtro_empresas", None)
            st.session_state.pop("filtro_obras", None)
            st.session_state.pop("filtro_fornecedores", None)
            st.rerun()

    modelo_atual = st.session_state.modelo_selecionado
    cfg_atual = MODELOS[modelo_atual]



    # =========================================================
    # 11. FILTROS E BUSCA
    # =========================================================

    st.markdown(
        '<div class="subtitulo-pagina">Filtros</div>',
        unsafe_allow_html=True,
    )

    try:
        with st.spinner("Carregando opções dos filtros..."):
            if modelo_atual == "empObr":
                opcoes_empresas = carregar_opcoes_empresas()

                empresas_map = dict(zip(
                    opcoes_empresas["Codigo_emp"],
                    opcoes_empresas["DescricaoEmpresa"],
                ))

                empresas_ids = list(empresas_map.keys())

                if "filtro_empresas" not in st.session_state:
                    st.session_state["filtro_empresas"] = []

                empresas_sel = st.multiselect(
                    "Empresa",
                    options=empresas_ids,
                    format_func=lambda codigo: empresas_map[codigo],
                    placeholder="Todas as empresas",
                    key="filtro_empresas",
                )

                opcoes_obras = carregar_opcoes_obras(
                    tuple(empresas_sel)
                )

                obras_map = dict(zip(
                    opcoes_obras["cod_obr"],
                    opcoes_obras["DescricaoObras"],
                ))

                obras_ids = list(obras_map.keys())

                # Remove seleções que não pertencem às empresas atuais.
                obras_anteriores = st.session_state.get("filtro_obras", [])
                obras_validas = [
                    obra for obra in obras_anteriores
                    if obra in obras_ids
                ]

                if obras_validas != obras_anteriores:
                    st.session_state["filtro_obras"] = obras_validas

                obras_sel = st.multiselect(
                    "Obra",
                    options=obras_ids,
                    format_func=lambda codigo: obras_map[codigo],
                    placeholder="Todas as obras",
                    key="filtro_obras",
                )

                fornecedores_sel = []

            else:
                opcoes_fornecedores = carregar_opcoes_fornecedores()

                fornecedores_ids = (
                    opcoes_fornecedores["ChqNome_Pag"]
                    .astype(str)
                    .tolist()
                )

                fornecedores_sel = st.multiselect(
                    "Fornecedor",
                    options=fornecedores_ids,
                    placeholder="Todos os fornecedores",
                    key="filtro_fornecedores",
                )

                empresas_sel = []
                obras_sel = []

    except Exception:
        logger.exception("Erro ao carregar as opções dos filtros.")
        st.error("Não foi possível carregar as opções dos filtros.")
        st.stop()

    st.caption(
        "Deixe um filtro sem seleção para considerar todos os valores."
    )


    buscar = st.button(
        "🔎 Buscar dados",
        type="primary",
        use_container_width=True,
        key="buscar_dados",
    )

    if buscar:
        # Guarda os filtros confirmados, sem executar buscas ao alterá-los.
        st.session_state["filtros_confirmados"] = {
            "modelo": modelo_atual,
            "empresas": tuple(empresas_sel),
            "obras": tuple(obras_sel),
            "fornecedores": tuple(fornecedores_sel),
        }

        st.session_state["busca_executada"] = True

    filtros = st.session_state.get("filtros_confirmados")

    if (
        not st.session_state.get("busca_executada", False)
        or not filtros
    ):
        st.info(
            "Selecione os filtros e clique em 'Buscar dados' "
            "para carregar os pagamentos."
        )


    if st.session_state.get("busca_executada", False) and filtros:
        filtros = st.session_state.get("filtros_confirmados", {})

        empresas_confirmadas = filtros.get("empresas", ())
        obras_confirmadas = filtros.get("obras", ())
        fornecedores_confirmados = filtros.get("fornecedores", ())

        try:
            df = carregar_dados(
                empresas=empresas_confirmadas,
                obras=obras_confirmadas,
                fornecedores=fornecedores_confirmados
            )

            if df.empty:
                st.warning("Nenhum pagamento encontrado para os filtros selecionados.")
                st.stop()

        except Exception as erro:
            st.error(f"Não foi possível consultar os pagamentos: {erro}")
            logging.exception("Falha ao carregar os pagamentos")
            st.stop()

        if df.empty:
            st.warning(
                "Nenhum pagamento encontrado. Altere os filtros "
                "e clique em 'Buscar dados' novamente."
            )
            st.stop()

        df["TotalRecebido"] = pd.to_numeric(
            df["TotalRecebido"],
            errors="coerce",
        ).fillna(0)

        df["TotalAno"] = pd.to_numeric(
            df["TotalAno"],
            errors="coerce",
        ).fillna(0)

        for coluna in [
            "DescricaoEmpresa",
            "DescricaoObras",
            "ChqNome_Pag",
        ]:
            df[coluna] = df[coluna].where(df[coluna].notna(), None)

        data_inicio = pd.Timestamp(df["DataInicio"].iloc[0])
        data_fim = pd.Timestamp(df["DataFim"].iloc[0])

        periodo = (
            f"Período: {data_inicio:%d/%m/%Y} "
            f"a {data_fim:%d/%m/%Y}"
        )

        df_filtrado = df.copy()

        st.caption(
            f"{len(df_filtrado):,} registros retornados."
            .replace(",", ".")
        )

        # =========================================================
        # 12. RESUMO E INDICADORES
        # =========================================================

        st.markdown(
            '<div class="subtitulo-pagina">Resumo da análise</div>',
            unsafe_allow_html=True,
        )

        metricas = calcular_metricas(df, df_filtrado)

        col_geral, col_filtrado = st.columns(2)

        with col_geral:
            st.markdown(
                f"""
                <div class="card-metrica destaque">
                    <div class="rotulo">Total geral do período</div>
                    <div class="valor">{brl(metricas["total_geral"])}</div>
                    <div class="detalhe">{html.escape(periodo)}</div>
                </div>
                """,
                unsafe_allow_html=True,
            )

        with col_filtrado:
            st.markdown(
                f"""
                <div class="card-metrica">
                    <div class="rotulo">Total dos filtros aplicados</div>
                    <div class="valor">{brl(metricas["total_filtrado"])}</div>
                    <div class="detalhe">Soma dos pagamentos correspondentes à seleção</div>
                </div>
                """,
                unsafe_allow_html=True,
            )

        col_emp, col_obr, col_forn = st.columns(3)

        indicadores = [
            (col_emp, "Empresas comparadas", metricas["empresas"]),
            (col_obr, "Obras comparadas", metricas["obras"]),
            (col_forn, "Fornecedores analisados", metricas["fornecedores"]),
        ]

        for coluna, rotulo, valor in indicadores:
            with coluna:
                st.markdown(
                    f"""
                    <div class="card-metrica">
                        <div class="rotulo">{rotulo}</div>
                        <div class="valor">{valor:,}</div>
                        <div class="detalhe">Conforme os filtros atuais</div>
                    </div>
                    """.replace(",", "."),
                    unsafe_allow_html=True,
                )


        # =========================================================
        # 13. RESULTADOS CONSOLIDADOS
        # =========================================================

        st.markdown(
            '<div class="subtitulo-pagina">Resultados consolidados</div>',
            unsafe_allow_html=True,
        )

        if df_filtrado.empty:
            st.info(
                "Nenhum registro corresponde aos filtros selecionados. "
                "Altere os filtros para visualizar resultados."
            )
            tabela_modelo = pd.DataFrame(
                columns=(
                    ["Empresa", "Obra", "Fornecedor", "Total Pago"]
                    if modelo_atual == "empObr"
                    else ["Fornecedor", "Obras", "Total Pago"]
                )
            )

        else:
            inicio_agrupamento = time.perf_counter()
            tabela_modelo = montar_modelo(df_filtrado, modelo_atual)

            logger.info(
                "Agrupamento do modelo concluído em %.2f segundos.",
                time.perf_counter() - inicio_agrupamento,
            )

            st.caption(
                f"Modelo: {cfg_atual['nome']} · "
                f"{len(tabela_modelo):,} registros · {periodo}"
                .replace(",", ".")
            )

            st.dataframe(
                formatar_colunas_monetarias(tabela_modelo),
                use_container_width=True,
                hide_index=True,
                height=360,
            )

        st.session_state["dados_ranking"] = df_filtrado.copy()
        st.session_state["tabela_ranking"] = tabela_modelo.copy()
        st.session_state["periodo_ranking"] = periodo
        st.session_state["modelo_ranking"] = modelo_atual


with aba_relatorios:
    tabela_modelo = st.session_state.get("tabela_ranking")
    periodo = st.session_state.get("periodo_ranking")
    modelo_atual = st.session_state.get("modelo_ranking")
    if tabela_modelo is None or periodo is None or modelo_atual not in MODELOS:
        st.info("Faça uma busca na aba Ranking para visualizar os relatórios.")
    else:
        cfg_atual = MODELOS[modelo_atual]
        # =========================================================
        # 14. PRÉVIAS DOS ARQUIVOS
        # =========================================================

        st.markdown(
            '<div class="subtitulo-pagina">Prévia dos relatórios</div>',
            unsafe_allow_html=True,
        )

        col_excel, col_pdf = st.columns(2)

        previa_excel = formatar_colunas_monetarias(
            tabela_modelo.head(LIMITE_PREVIA)
        )

        with col_excel:
            with st.container(border=True):
                st.markdown(
                    '<div class="modelo-titulo">📗 Excel (.xlsx)</div>',
                    unsafe_allow_html=True,
                )
                st.caption(cfg_atual["nome"])
                st.caption(periodo)

                st.dataframe(
                    previa_excel,
                    hide_index=True,
                    use_container_width=True,
                    height=215,
                )

        with col_pdf:
            with st.container(border=True):
                st.markdown(
                    '<div class="modelo-titulo">📕 PDF (.pdf)</div>',
                    unsafe_allow_html=True,
                )

                html_previa = previa_pdf_html(
                    tabela_modelo,
                    modelo_atual,
                    periodo,
                )

                components.html(
                    html_previa,
                    height=300,
                    scrolling=True,
                )


        # =========================================================
        # 15. DOWNLOADS SOB DEMANDA
        # =========================================================

        assinatura_atual = assinatura_dataframe(tabela_modelo)

        st.markdown(
            '<div class="subtitulo-pagina">Baixar relatórios</div>',
            unsafe_allow_html=True,
        )

        if tabela_modelo.empty:
            st.info(
                "Os downloads estarão disponíveis quando houver resultados "
                "para o modelo e os filtros escolhidos."
            )

        else:
            col_download_excel, col_download_pdf = st.columns(2)

            with col_download_excel:
                if st.button(
                    "Preparar Excel",
                    use_container_width=True,
                    key="preparar_excel",
                ):
                    inicio = time.perf_counter()
                    try:
                        st.session_state["arquivo_excel"] = gerar_excel(tabela_modelo)
                        st.session_state["chave_excel"] = (
                            modelo_atual,
                            assinatura_dataframe(tabela_modelo),
                        )
                        logger.info(
                            "Preparação do Excel concluída em %.2f segundos.",
                            time.perf_counter() - inicio,
                        )
                    except Exception:
                        logger.exception("Erro ao preparar o Excel.")
                        st.error("Não foi possível gerar o arquivo Excel.")

            with col_download_pdf:
                if st.button(
                    "Preparar PDF",
                    use_container_width=True,
                    key="preparar_pdf",
                ):
                    inicio = time.perf_counter()
                    try:
                        st.session_state["arquivo_pdf"] = gerar_pdf(
                            tabela_modelo,
                            cfg_atual["nome"],
                            periodo,
                        )
                        st.session_state["chave_pdf"] = (
                            modelo_atual,
                            assinatura_dataframe(tabela_modelo),
                            periodo,
                        )
                        logger.info(
                            "Preparação do PDF concluída em %.2f segundos.",
                            time.perf_counter() - inicio,
                        )
                    except Exception:
                        logger.exception("Erro ao preparar o PDF.")
                        st.error("Não foi possível gerar o arquivo PDF.")

            chave_atual_excel = (
                modelo_atual,
                assinatura_atual,
            )

            chave_atual_pdf = (
                modelo_atual,
                assinatura_atual,
                periodo,
            )

            # Não permite baixar um arquivo antigo após mudar os filtros.
            if st.session_state.get("chave_excel") != chave_atual_excel:
                st.session_state.pop("arquivo_excel", None)
                st.session_state.pop("chave_excel", None)

            if st.session_state.get("chave_pdf") != chave_atual_pdf:
                st.session_state.pop("arquivo_pdf", None)
                st.session_state.pop("chave_pdf", None)

            if "arquivo_excel" in st.session_state:
                st.download_button(
                    "⬇ Baixar relatório Excel",
                    data=st.session_state["arquivo_excel"],
                    file_name=f"{cfg_atual['arquivo']}.xlsx",
                    mime=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                    use_container_width=True,
                    key="download_excel",
                )

            if "arquivo_pdf" in st.session_state:
                st.download_button(
                    "⬇ Baixar relatório PDF",
                    data=st.session_state["arquivo_pdf"],
                    file_name=f"{cfg_atual['arquivo']}.pdf",
                    mime="application/pdf",
                    use_container_width=True,
                    key="download_pdf",
                )

        st.caption(
            "Os dados são consultados no UAU e armazenados temporariamente "
            "em cache por até cinco minutos. Os tempos técnicos são registrados "
            "nos logs do servidor."
        )

with aba_abc:
    dados_abc = st.session_state.get("dados_ranking")
    if dados_abc is None:
        st.info("Faça uma busca na aba Ranking para visualizar a Curva ABC.")
    else:
        _titulo_secao("Análise ABC")
        exibir_grafico_abc(dados_abc)
